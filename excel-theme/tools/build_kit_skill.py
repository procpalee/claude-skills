# -*- coding: utf-8 -*-
"""
build_kit_skill.py - 강의·가이드 킷용 excel-theme 스킬 빌드 (default 테마 전용, openpyxl + Office.js 공용).

    python -X utf8 tools/build_kit_skill.py   →  dist/excel-theme-kit/excel-theme/ + excel-theme.zip

문서 원본은 kit/(SKILL.md · references/ · scripts/sync_theme.py), 코드는 이 스킬 폴더의 원본을 그대로 가져온다.
검증(패리티·Node 목·문서 예제·theme_lint·커스터마이징 모의)이 하나라도 실패하면 ZIP 을 만들지 않는다.
자세한 내용은 kit/README.md.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

MASTER = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # excel-theme
SRC = os.path.join(MASTER, "kit")
DIST = os.path.join(MASTER, "dist", "excel-theme-kit")
WORK = os.path.join(DIST, "_work")                                     # 테스트·모의 작업 폴더
OUT = os.path.join(DIST, "excel-theme")
ZIP = os.path.join(DIST, "excel-theme.zip")
ENV = dict(os.environ, PYTHONIOENCODING="utf-8")


def read(p):
    with open(p, encoding="utf-8") as f:
        return f.read().replace("\r\n", "\n")


def write(p, t):
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8", newline="\n") as f:
        f.write(t)


def block_end(t, i):
    """t[i] 이후 첫 '{' 부터 짝이 맞는 '}' 위치 다음."""
    j = t.index("{", i)
    d = 0
    for k in range(j, len(t)):
        if t[k] == "{":
            d += 1
        elif t[k] == "}":
            d -= 1
            if d == 0:
                return k + 1
    raise ValueError("brace")


def default_only(t):
    """PALETTES 를 default 하나로, ALIASES 를 default 로 가는 것만 남긴다."""
    i = t.index("PALETTES = {")
    end = block_end(t, i)
    d0 = t.index('    "default": {', i)
    d1 = block_end(t, d0)
    pal = ("PALETTES = {\n"
           "    # 킷 버전: 테마는 default 하나(정산표·결산 회색 표). 색을 바꾸려면 아래 값만 고치고\n"
           "    # python scripts/sync_theme.py 를 돌린다(references/customize.md).\n"
           + t[d0:d1] + ",\n}")
    t = t[:i] + pal + t[end:]
    a = t.index("ALIASES = {")
    t = t[:a] + 'ALIASES = {"closing": "default"}' + t[block_end(t, a):]
    t = re.sub(r"# 미지정 시 기본 테마\. 공식 카탈로그 3종.*?(?=DEFAULT_THEME = )",
               "# 기본 테마(킷 버전은 default 하나).\n", t, flags=re.S)
    t = t.replace('DEFAULT_THEME = os.environ.get("EXCEL_THEME_DEFAULT", "default")', 'DEFAULT_THEME = "default"')
    assert 'PALETTES = {\n' in t and '"audit": {' not in t
    return t


def run(args):
    p = subprocess.run([sys.executable] + args, capture_output=True, text=True, encoding="utf-8", env=ENV)
    if p.returncode != 0:
        raise SystemExit("실패: %s\n%s\n%s" % (" ".join(args), p.stdout[-2500:], p.stderr[-2500:]))
    return p.stdout


def patched_test(name, sd):
    """officejs/tests 의 검사 스크립트를 복사해 대상 경로만 패키지 scripts/ 로 바꾼다."""
    t = read(os.path.join(MASTER, "officejs", "tests", name))
    t = re.sub(r"^SKILL = .*$", lambda m: 'SKILL = r"%s"' % sd, t, count=1, flags=re.M)
    t = re.sub(r"^ROOT = .*$", lambda m: 'ROOT = r"%s"' % sd, t, count=1, flags=re.M)
    dst = os.path.join(WORK, "tests", name)
    write(dst, t)
    return dst


# references/openpyxl.md 예제와 같은 호출 - 문서가 실제 함수와 맞는지 확인
DOC_SMOKE = r'''
import sys; sys.path.insert(0, r"%(sd)s")
from openpyxl import Workbook, load_workbook
from excel_theme import (write_table, write_title, write_caption, write_section_bar, set_tab,
                         mark_cells, apply_theme, style_total_row, get_theme, THEMES)
assert list(THEMES) == ["default"], THEMES
wb = Workbook(); ws = wb.active; ws.title = "재무상태표"
rows = [["유동자산", 1200, 1100], ["비유동자산", 800, 750], ["자산총계", "=SUM(C5:C6)", "=SUM(D5:D6)"]]
write_table(ws, ["계정", "당기", "전기"], rows, title="재무상태표", currency_cols=[1, 2], total_rows=[3])
set_tab(ws, "output")
ws2 = wb.create_sheet("20 연령분석")
assert write_title(ws2, last_col=7, tab="calc") == 4
write_caption(ws2, "B12", "표1. 연령별 잔액")
write_section_bar(ws2, 6, "1. 총괄표", last_col=8)
mark_cells(ws2, get_theme(), input=["C6:D7"], linked=["C10"])
ws3 = wb.create_sheet("기존표")
for r, row in enumerate([["항목", "금액"], ["A", 1], ["B", 2], ["합계", "=SUM(C5:C6)"]], start=4):
    for c, v in enumerate(row, start=2):
        ws3.cell(r, c, v)
apply_theme(ws3, header_row=4, data_range="B4:C7", title_cell="B2", number_format_cols={"C": "accounting"})
style_total_row(ws3, get_theme(), row=7, min_col=2, max_col=3)
out = r"%(out)s"; wb.save(out)
print("doc-smoke OK header", load_workbook(out)["재무상태표"]["B4"].fill.fgColor.rgb)
'''


def test_package(root, label, expect_header):
    sd = os.path.join(root, "scripts")
    print(" [%s] %s" % (label, run([os.path.join(sd, "sync_theme.py"), "--check"]).strip()))
    for t in ("check_parity.py", "check_officejs.py"):
        print(" [%s] %-18s %s" % (label, t, run([patched_test(t, sd)]).strip().splitlines()[-1]))
    xl = os.path.join(WORK, "tests", "doc_smoke_%s.xlsx" % label)
    py = os.path.join(WORK, "tests", "doc_smoke.py")
    write(py, DOC_SMOKE % {"sd": sd, "out": xl})
    o = run([py]).strip()
    print(" [%s] %s" % (label, o))
    assert o.endswith(expect_header), o
    print(" [%s] %s" % (label, run([os.path.join(sd, "theme_lint.py"), xl, "default"]).strip().splitlines()[-1]))
    js = run([os.path.join(sd, "make_code.py"), "table", json.dumps({"n_cols": 2, "n_rows": 2})])
    assert "Excel.run" in js


def main():
    shutil.rmtree(DIST, ignore_errors=True)
    S = os.path.join(OUT, "scripts")
    R = os.path.join(OUT, "references")
    write(os.path.join(OUT, "SKILL.md"), read(os.path.join(SRC, "SKILL.md")))
    write(os.path.join(S, "excel_theme.py"), default_only(read(os.path.join(MASTER, "excel_theme.py"))))
    for fn in ("theme_plan.py", "plan.py", "theme_lint.py"):
        write(os.path.join(S, fn), read(os.path.join(MASTER, fn)))
    for fn in ("officejs_plan.py", "make_code.py"):
        write(os.path.join(S, fn), read(os.path.join(MASTER, "officejs", fn)))
    write(os.path.join(S, "sync_theme.py"), read(os.path.join(SRC, "scripts", "sync_theme.py")))
    for fn in ("openpyxl.md", "officejs.md", "customize.md", "rules.md"):
        write(os.path.join(R, fn), read(os.path.join(SRC, "references", fn)))
    print(run([os.path.join(S, "sync_theme.py")]).strip())
    data = json.load(open(os.path.join(S, "theme.json"), encoding="utf-8"))
    assert list(data["themes"]) == ["default"], list(data["themes"])
    m = re.match(r"---\n(.*?)\n---\n", read(os.path.join(OUT, "SKILL.md")), re.S)
    fm = dict(l.split(":", 1) for l in m.group(1).splitlines() if ":" in l)
    assert fm["name"].strip() == "excel-theme" and 0 < len(fm["description"].strip()) <= 1024
    test_package(OUT, "원본", expect_header="F2F2F2")

    # 실습 I 모의: 색만 바꾼 사본에서 두 엔진이 같이 바뀌고, 동기화 누락을 --check 가 잡는지
    sim = os.path.join(WORK, "sim", "my-excel-theme")
    shutil.copytree(OUT, sim)
    p = os.path.join(sim, "scripts", "excel_theme.py")
    t = read(p)
    for k, v in (("excel_header_fill", "1F3864"), ("excel_header_text", "FFFFFF"),
                 ("excel_total_fill", "D9E1F2"), ("border_in", "D9D9D9")):
        t, n = re.subn(r'("%s":\s*")[0-9A-Fa-f]{6}(")' % k, r"\g<1>%s\2" % v, t)
        assert n == 1, k
    write(p, t)
    stale = subprocess.run([sys.executable, os.path.join(sim, "scripts", "sync_theme.py"), "--check"],
                           capture_output=True, env=ENV)
    assert stale.returncode == 1, "동기화 누락을 --check 가 못 잡음"
    run([os.path.join(sim, "scripts", "sync_theme.py")])
    js = run([os.path.join(sim, "scripts", "make_code.py"), "table",
              json.dumps({"n_cols": 3, "n_rows": 3, "total_rows": [3]})]).upper()
    assert "#1F3864" in js and "#D9E1F2" in js, "Office.js 코드에 새 색 없음"
    test_package(sim, "커스텀", expect_header="1F3864")

    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for root, dirs, files in os.walk(OUT):
            dirs[:] = sorted(d for d in dirs if d != "__pycache__")
            for fn in sorted(files):
                full = os.path.join(root, fn)
                z.write(full, os.path.relpath(full, DIST).replace(os.sep, "/"))
    names = zipfile.ZipFile(ZIP).namelist()
    print("OK  %s  (%d files, %.1f KB)" % (ZIP, len(names), os.path.getsize(ZIP) / 1024))
    for n in names:
        print("    " + n)


if __name__ == "__main__":
    main()
