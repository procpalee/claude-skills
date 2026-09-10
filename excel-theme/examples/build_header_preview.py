# -*- coding: utf-8 -*-
"""계정조서 헤더 템플릿 미리보기 — grid / line / block 3종 (excel-theme write_header)."""
import sys
sys.path.insert(0, r"C:\Users\PC\.claude\skills\excel-theme")
import openpyxl
from excel_theme import write_header, write_table, write_caption, fit_rows, HEADER_STYLES
OUT = sys.argv[1] if len(sys.argv) > 1 else "조서헤더_템플릿_미리보기.xlsx"
META = {"회사명": "흥구석유(주)", "조서번호": "6001", "기준일": "2026.06.30", "작성자": "이재현", "작성일자": "2026-09-04",
        "검토자": "박승진", "검토일자": "2026-09-05", "상태": "검토 중"}
DESC = {"grid": "A. grid — 워드 조서와 동일: 라벨음영 격자 정보표(2행)가 제목 위, 제목 밴드 아래에 본문 (권장)",
        "line": "B. line — 제목 밴드 + 한 줄 메타(회색 9pt, 하단선). 가장 간결, 표 위 공간 최소",
        "block": "C. block — 제목 밴드 + 좌(회사·조서·기준일)/우(작성·검토·상태) 2단 정보블록. 메모형"}
wb = openpyxl.Workbook(); wb.remove(wb.active)
for st in HEADER_STYLES:
    ws = wb.create_sheet(f"{st}")
    hdr = write_header(ws, META, style=st, last_col=7)   # 제목 밴드 = 시트명 (확정 2026-09-10)
    last = write_table(ws, ["항목", "내용", "담당", "문서", "통제", "주장"], [
        ["1 주문접수", "거래처 주문 접수 → 유류수탁처리부 기록 → 수송부 전달", "영업부", "유류수탁처리부", "K1 주문 승인", "E/O"],
        ["2 정유사 주문", "GS ePartners 주문 입력", "수송부", "GS ePartners", "주문 = 지시", "C"],
        ["3 출하 · 납품", "납품 시 출하명세표(인수증) 작성, 주유소 날인", "운송기사", "인수증", "K2 거래처 날인", "E/O"],
    ], theme="default", header_row=hdr)
    write_caption(ws, f"B{hdr+6}", "1. 다음 표 (캡션 예시)")
    from openpyxl.styles import Font
    ws[f"B{hdr+8}"] = DESC[st]
    ws[f"B{hdr+8}"].font = Font(name="맑은 고딕", size=10, color="595959")
    for col, w in {"A": 2, "B": 18, "C": 44, "D": 12, "E": 22, "F": 22, "G": 12}.items():
        ws.column_dimensions[col].width = w
    fit_rows(ws, min_col=2, max_col=7)
wb.save(OUT); print("saved", OUT)
